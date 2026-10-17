"""Prepare a course inventory by copying stored workbook hyperlink targets.

Paths are explicit and relative to the caller's working directory. Importing
this module does not search for a checkout, open files, or perform extraction.
Only ordinary .xlsx workbooks are supported; macros are outside this contract.
"""

from __future__ import annotations

import argparse
import os
import sys
from collections.abc import Sequence
from dataclasses import dataclass
from pathlib import Path
from tempfile import NamedTemporaryFile

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet


class ExtractionError(ValueError):
    """Indicate invalid extraction arguments or unsupported workbook input."""


@dataclass(frozen=True, slots=True)
class ExtractionResult:
    """Describe a successfully published processed workbook.

    copied_count counts cells assigned a stored hyperlink target, including
    cells that already contained that same value. rows_examined includes
    worksheet rows 2 through max_row, including any blank/formatted rows.
    """

    source_path: Path
    destination_path: Path
    sheet_name: str
    copied_count: int
    rows_examined: int
    url_column_created: bool

    @property
    def skipped_count(self) -> int:
        """Return the number of examined rows left unchanged."""
        return self.rows_examined - self.copied_count


def _path(value: str | os.PathLike[str], label: str) -> Path:
    """Resolve an explicit path using the native home and working directory."""
    try:
        raw = os.fspath(value)
        if not isinstance(raw, str) or not raw.strip() or "\x00" in raw:
            raise ValueError
        path = Path(raw).expanduser().resolve()
    except (OSError, RuntimeError, TypeError, ValueError):
        raise ExtractionError(f"{label} must be a usable file path.") from None
    if path.suffix.lower() != ".xlsx":
        raise ExtractionError(f"{label} must name an .xlsx workbook.")
    return path


def _name(value: str, label: str) -> str:
    """Require a nonblank configured worksheet or column name."""
    if not isinstance(value, str) or not value.strip():
        raise ExtractionError(f"{label} must be a nonblank string.")
    return value


def _check_paths(source: Path, destination: Path, overwrite: bool) -> None:
    """Refuse missing input, source aliases, and unapproved replacement."""
    if source == destination:
        raise ExtractionError("Source and destination must be different files.")
    if not source.exists():
        raise FileNotFoundError("Source workbook does not exist.")
    if not source.is_file():
        raise ExtractionError("Source workbook must be a regular file.")
    if destination.exists():
        if source.samefile(destination):
            raise ExtractionError(
                "Source and destination refer to the same file."
            )
        if not destination.is_file():
            raise ExtractionError("Destination must name a workbook file.")
        if not overwrite:
            raise FileExistsError(
                "Destination exists; use overwrite=True or --overwrite "
                "to replace it."
            )


def _columns(
    worksheet: Worksheet, name_column: str, url_column: str
) -> tuple[int, int, bool]:
    """Locate normalized headers and append the URL column when absent."""
    columns: dict[str, int] = {}
    for cell in worksheet[1]:
        value = cell.value
        if value is None or isinstance(value, str) and not value.strip():
            continue
        if not isinstance(value, str):
            raise ExtractionError("Workbook headers must contain text.")
        key = value.strip().casefold()
        if key in columns:
            raise ExtractionError(
                "Duplicate normalized workbook headers at columns "
                f"{columns[key]} and {cell.column}."
            )
        columns[key] = cell.column
    name_index = columns.get(name_column.strip().casefold())
    if name_index is None:
        raise ExtractionError("Missing required course-name column.")
    url_index = columns.get(url_column.strip().casefold())
    created = url_index is None
    if url_index is None:
        if worksheet.max_column >= 16_384:
            raise ExtractionError("Worksheet has no room for a URL column.")
        url_index = worksheet.max_column + 1
        worksheet.cell(row=1, column=url_index, value=url_column.strip())
    return name_index, url_index, created


def _copy_targets(
    worksheet: Worksheet, name_index: int, url_index: int
) -> tuple[int, int]:
    """Copy actual external targets as text, preserving rows without one."""
    copied = 0
    rows_examined = max(worksheet.max_row - 1, 0)
    for row in range(2, worksheet.max_row + 1):
        hyperlink = worksheet.cell(row=row, column=name_index).hyperlink
        target = None if hyperlink is None else hyperlink.target
        if (
            not isinstance(target, str)
            or not target.strip()
            or target.lstrip().startswith("#")
        ):
            continue
        output = worksheet.cell(row=row, column=url_index)
        if isinstance(output, MergedCell):
            raise ExtractionError("Unmerge URL cells before extraction.")
        output.value = target
        # A stored target is data even if it starts with a formula marker.
        output.data_type = "s"
        copied += 1
    return copied, rows_examined


def _publish(workbook: Workbook, destination: Path, overwrite: bool) -> None:
    """Publish a fully saved workbook while preserving existing files."""
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with NamedTemporaryFile(
            dir=destination.parent, prefix=".d2l-", suffix=".xlsx", delete=False
        ) as temporary:
            temporary_path = Path(temporary.name)
            try:
                workbook.save(temporary)
            except OSError:
                raise OSError("Cannot save the processed workbook.") from None
            except Exception:
                message = "Cannot serialize the processed workbook."
                raise OSError(message) from None
            temporary.flush()
            os.fsync(temporary.fileno())
        if overwrite:
            os.replace(temporary_path, destination)
        else:
            # Linking publishes the complete file without replacing a raced
            # destination. Unsupported filesystems fail without overwriting.
            os.link(temporary_path, destination)
    finally:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)


def extract_course_urls(
    source: str | os.PathLike[str],
    destination: str | os.PathLike[str],
    *,
    sheet_name: str = "D2L",
    name_column: str = "D2l_course_name",
    url_column: str = "CourseURL",
    overwrite: bool = False,
) -> ExtractionResult:
    """Copy external course-name hyperlink targets to a processed workbook.

    Args:
        source: Existing .xlsx file. Never modified by this function.
        destination: Explicit .xlsx output path. Its parents are created
            after workbook validation. Relative paths use the current CWD.
        sheet_name: Exact worksheet name to process.
        name_column: Course-name header, matched after trim and casefold.
        url_column: URL header with the same normalization. Added at the
            worksheet's end when absent.
        overwrite: Permit replacement of an existing destination workbook.
            Source/destination identity is refused even when true.

    Returns:
        Immutable result with resolved paths, copy/row counts, and whether
        the URL column was created. This replaces the former integer return.

    Raises:
        ExtractionError: For invalid arguments, unreadable workbook content,
            missing/duplicate headers, missing sheet, or aliased paths.
        FileNotFoundError: If the source does not exist.
        FileExistsError: If output exists without overwrite permission,
            including output created concurrently during extraction.
        OSError: If input access, saving, publication, or cleanup fails.

    Existing URL cells are unchanged without an external target. Internal
    locations and HYPERLINK formulas are not evaluated or parsed for links.
    Other data, formulas, sheets, links, and styles supported by openpyxl
    are preserved; unsupported Excel features are not guaranteed.

    Output is saved to a temporary file in the destination directory first.
    Default publication requires filesystem hard-link support and fails
    without replacing output if unavailable. Explicit overwrite uses atomic
    replacement. Orderly cleanup removes the temporary file.
    """
    if not isinstance(overwrite, bool):
        raise ExtractionError("overwrite must be a boolean.")
    sheet_name = _name(sheet_name, "sheet_name")
    name_column = _name(name_column, "name_column")
    url_column = _name(url_column, "url_column")
    if name_column.strip().casefold() == url_column.strip().casefold():
        raise ExtractionError("Name and URL columns must be different.")
    source_path = _path(source, "source")
    destination_path = _path(destination, "destination")
    _check_paths(source_path, destination_path, overwrite)
    with source_path.open("rb") as stream:
        try:
            workbook = load_workbook(stream, data_only=False, keep_links=True)
        except Exception:
            message = "Cannot parse the source XLSX workbook."
            raise ExtractionError(message) from None
    try:
        if sheet_name not in workbook.sheetnames:
            raise ExtractionError("Configured worksheet is missing.")
        worksheet = workbook[sheet_name]
        name_index, url_index, created = _columns(
            worksheet, name_column, url_column
        )
        copied, rows_examined = _copy_targets(worksheet, name_index, url_index)
        _publish(workbook, destination_path, overwrite)
    finally:
        workbook.close()
    return ExtractionResult(
        source_path=source_path,
        destination_path=destination_path,
        sheet_name=sheet_name,
        copied_count=copied,
        rows_examined=rows_examined,
        url_column_created=created,
    )


def main(argv: Sequence[str] | None = None) -> int:
    """Run extraction from explicit arguments and return an exit code."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path, help="Source .xlsx workbook")
    parser.add_argument("destination", type=Path, help="Processed .xlsx path")
    parser.add_argument("--sheet-name", default="D2L")
    parser.add_argument("--name-column", default="D2l_course_name")
    parser.add_argument("--url-column", default="CourseURL")
    parser.add_argument("--overwrite", action="store_true")
    arguments = parser.parse_args(argv)
    try:
        result = extract_course_urls(
            arguments.source,
            arguments.destination,
            sheet_name=arguments.sheet_name,
            name_column=arguments.name_column,
            url_column=arguments.url_column,
            overwrite=arguments.overwrite,
        )
    except (ExtractionError, FileNotFoundError, FileExistsError) as error:
        print(f"Extraction failed: {error}", file=sys.stderr)
        return 2
    except OSError as error:
        print(f"Extraction I/O failed: {error}", file=sys.stderr)
        return 1
    print(
        f"Saved {result.copied_count} course URLs to {result.destination_path}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
