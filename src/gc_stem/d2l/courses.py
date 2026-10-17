"""Load course inventories without opening Brightspace or changing files.

Excel uses the configured worksheet and column names. CSV requires the
configured ID header; display-name and course-URL columns are optional.
Only HTTPS course-home URLs on the configured tenant are accepted.
"""

from __future__ import annotations

import csv
import hashlib
import ipaddress
import math
import re
from collections.abc import Generator, Sequence
from contextlib import closing
from dataclasses import dataclass, replace
from io import BytesIO, StringIO
from pathlib import Path
from typing import Literal
from urllib.parse import urlsplit

from openpyxl import load_workbook

from .config import CourseFormat, CourseSettings, Settings
from .models import CourseRef


DiagnosticCode = Literal[
    "INVALID_ROW", "DUPLICATE_COURSE", "NAME_DIFFERENCE"
]
_HOME_ROUTE = re.compile(r"(/d2l/(?:p/)?home)/([0-9]+)/?")
# Excel numeric cells retain at most 15 significant decimal digits.
_MAX_NUMERIC_ID = 10**15 - 1


@dataclass(frozen=True, slots=True)
class InventoryDiagnostic:
    """Describe a row problem without retaining rejected cell contents.

    row is the worksheet row or the physical starting line of a CSV
    record. related_row identifies the first occurrence of a duplicate.
    INVALID_ROW is fatal; the other codes are nonfatal warnings.
    """

    code: DiagnosticCode
    row: int
    message: str
    related_row: int | None = None
    org_unit_id: str | None = None


class InventoryError(ValueError):
    """Indicate unreadable input, invalid structure, or invalid course rows."""

    def __init__(
        self,
        message: str,
        *,
        diagnostics: tuple[InventoryDiagnostic, ...] = (),
    ) -> None:
        """Attach structured row diagnostics to an actionable error."""
        super().__init__(message)
        self.diagnostics = diagnostics


@dataclass(frozen=True, slots=True)
class Inventory:
    """Hold validated courses and metadata for the exact input bytes read.

    row_count includes all nonblank data rows, including duplicates.
    Course order and display metadata come from each identity's first row.
    Raw workbook/CSV contents are not retained in the result.
    """

    courses: tuple[CourseRef, ...]
    source_path: Path
    source_format: CourseFormat
    sheet_name: str | None
    sha256: str
    row_count: int
    diagnostics: tuple[InventoryDiagnostic, ...] = ()

    @property
    def duplicate_count(self) -> int:
        """Return the number of additional rows for existing identities."""
        return self.row_count - len(self.courses)


def _blank(value: object) -> bool:
    """Recognize absent cells and whitespace-only text."""
    return value is None or isinstance(value, str) and not value.strip()


def _course_id(value: object) -> str:
    """Normalize an unambiguous positive decimal ID without guessing."""
    if isinstance(value, bool):
        raise InventoryError("Course ID must not be a boolean.")
    if isinstance(value, int):
        if 0 < value <= _MAX_NUMERIC_ID:
            return str(value)
    elif isinstance(value, float):
        if (
            math.isfinite(value)
            and value.is_integer()
            and 0 < value <= _MAX_NUMERIC_ID
        ):
            return str(int(value))
    elif isinstance(value, str):
        text = value.strip()
        if text.isascii() and text.isdecimal():
            normalized = text.lstrip("0")
            if normalized:
                return normalized
    raise InventoryError(
        "Course ID must be a positive decimal integer; formulas, fractions, "
        "and scientific-notation strings are unsupported. Numeric IDs longer "
        "than 15 digits must be stored as text."
    )


def _url_identity(value: object, tenant_origin: str) -> tuple[str, str]:
    """Validate one recognized home route and return its ID and clean URL."""
    if not isinstance(value, str):
        raise InventoryError("Course URL must contain stored URL text.")
    raw = value.strip()
    if (
        any(character.isspace() or ord(character) < 32 for character in raw)
        or any(character in raw for character in "\\?#")
        or "\x7f" in raw
    ):
        raise InventoryError(
            "Course URL must not contain whitespace, queries, or fragments."
        )
    try:
        parts = urlsplit(raw)
        host = parts.hostname
        port = parts.port
    except ValueError:
        raise InventoryError("Course URL is malformed.") from None
    if (
        parts.scheme != "https"
        or not host
        or parts.username is not None
        or parts.password is not None
        or port == 0
    ):
        raise InventoryError(
            "Course URL must be an absolute HTTPS URL without credentials."
        )
    try:
        try:
            host = ipaddress.ip_address(host).compressed
        except ValueError:
            host = host.encode("idna").decode("ascii").lower()
    except UnicodeError:
        raise InventoryError("Course URL has an invalid hostname.") from None
    authority = f"[{host}]" if ":" in host else host
    if port not in (None, 443):
        authority = f"{authority}:{port}"
    if f"https://{authority}" != tenant_origin:
        raise InventoryError("Course URL is outside the configured tenant.")
    route = _HOME_ROUTE.fullmatch(parts.path)
    if route is None:
        raise InventoryError(
            "Course URL must use /d2l/home/ID or /d2l/p/home/ID."
        )
    identifier = _course_id(route[2])
    return identifier, f"{tenant_origin}{route[1]}/{identifier}"


def _csv_rows(
    data: bytes,
) -> Generator[tuple[int, tuple[object, ...]], None, None]:
    """Read strict UTF-8 CSV and retain each record's starting line."""
    try:
        text = data.decode("utf-8-sig")
    except UnicodeError:
        raise InventoryError("CSV inventory must use UTF-8.") from None
    with StringIO(text, newline="") as source:
        reader = csv.reader(source, strict=True)
        while True:
            row = reader.line_num + 1
            try:
                values = next(reader)
            except StopIteration:
                return
            except csv.Error:
                raise InventoryError(
                    f"Malformed CSV inventory near line {row}."
                ) from None
            yield row, tuple(values)


def _excel_rows(
    data: bytes, sheet_name: str
) -> Generator[tuple[int, tuple[object, ...]], None, None]:
    """Read stored cell values without evaluating formulas or external links."""
    with BytesIO(data) as source:
        try:
            workbook = load_workbook(
                source, read_only=True, data_only=False, keep_links=False
            )
        except Exception:
            # openpyxl and its XML/ZIP backends expose varied parse errors.
            raise InventoryError("Cannot parse the XLSX inventory.") from None
        try:
            if sheet_name not in workbook.sheetnames:
                raise InventoryError(
                    "Configured courses.sheet is missing from the workbook."
                )
            try:
                rows = workbook[sheet_name].iter_rows(values_only=True)
                for row, values in enumerate(rows, start=1):
                    yield row, values
            except Exception:
                raise InventoryError(
                    "Cannot read cells from the XLSX inventory."
                ) from None
        finally:
            workbook.close()


def _columns(
    headers: Sequence[object], settings: CourseSettings
) -> tuple[int | None, int | None, int | None]:
    """Validate normalized headers and locate configured inventory columns."""
    positions: dict[str, int] = {}
    for index, header in enumerate(headers):
        if _blank(header):
            continue
        if not isinstance(header, str):
            raise InventoryError("Inventory headers must contain text.")
        key = header.strip().casefold()
        if key in positions:
            raise InventoryError(
                "Duplicate normalized inventory headers at columns "
                f"{positions[key] + 1} and {index + 1}."
            )
        positions[key] = index
    if not positions:
        raise InventoryError("Inventory has no header row.")

    names = (settings.name_column, settings.url_column, settings.id_column)
    selected = [name.strip().casefold() for name in names if name is not None]
    if len(set(selected)) != len(selected):
        raise InventoryError("Inventory column mappings must be distinct.")
    name_index, url_index, id_index = (
        positions.get(name.strip().casefold()) if name is not None else None
        for name in names
    )
    if settings.format == "xlsx":
        if url_index is None:
            raise InventoryError(
                "Missing courses.url_column header (CourseURL by default). "
                "Prepare the workbook with extract_course_urls() and use "
                "its processed output."
            )
        if name_index is None:
            raise InventoryError("Missing courses.name_column header.")
    if settings.id_column is not None and id_index is None:
        raise InventoryError("Missing courses.id_column header.")
    if settings.format == "csv" and id_index is None:
        raise InventoryError("CSV inventory requires courses.id_column.")
    return name_index, url_index, id_index


def _cell(values: Sequence[object], index: int | None) -> object:
    """Return an optional cell, including absent trailing CSV values."""
    if index is None or index >= len(values):
        return None
    return values[index]


def _course(
    values: Sequence[object],
    columns: tuple[int | None, int | None, int | None],
    tenant_origin: str,
    row: int,
) -> CourseRef:
    """Validate one identity and preserve its display name and source row."""
    name_value, url_value, id_value = (
        _cell(values, index) for index in columns
    )
    identifier = None if _blank(id_value) else _course_id(id_value)
    if _blank(url_value):
        if identifier is None:
            raise InventoryError("Course row has neither an ID nor a URL.")
        url = f"{tenant_origin}/d2l/home/{identifier}"
    else:
        url_id, url = _url_identity(url_value, tenant_origin)
        if identifier is not None and identifier != url_id:
            raise InventoryError("Course ID conflicts with its course URL.")
        identifier = url_id
    if _blank(name_value):
        name = ""
    elif isinstance(name_value, str):
        name = name_value.strip()
    else:
        raise InventoryError("Course display name must contain text.")
    return CourseRef(
        tenant_origin=tenant_origin,
        org_unit_id=identifier,
        name=name,
        course_url=url,
        source_rows=(row,),
    )


def _deduplicate(
    courses: dict[tuple[str, str], CourseRef],
    course: CourseRef,
    diagnostics: list[InventoryDiagnostic],
) -> None:
    """Merge source rows while retaining first order and display metadata."""
    key = course.tenant_origin, course.org_unit_id
    first = courses.get(key)
    if first is None:
        courses[key] = course
        return
    row, first_row = course.source_rows[0], first.source_rows[0]
    diagnostics.append(
        InventoryDiagnostic(
            code="DUPLICATE_COURSE",
            row=row,
            related_row=first_row,
            org_unit_id=course.org_unit_id,
            message="Duplicate course; the first occurrence is retained.",
        )
    )
    if course.name != first.name:
        diagnostics.append(
            InventoryDiagnostic(
                code="NAME_DIFFERENCE",
                row=row,
                related_row=first_row,
                org_unit_id=course.org_unit_id,
                message="Display name differs; the first name is retained.",
            )
        )
    courses[key] = replace(
        first, source_rows=first.source_rows + course.source_rows
    )


def load_courses(settings: Settings) -> Inventory:
    """Read and validate the entire inventory before returning any courses.

    Args:
        settings: Validated settings returned by config.load_settings().

    Returns:
        Unique courses in first-input order, source metadata, SHA-256 of
        the bytes parsed, and nonfatal duplicate/display-name diagnostics.

    Raises:
        InventoryError: For unreadable files, invalid structure, no course
            rows, or invalid identities. Row failures are collected in
            exception.diagnostics; the message previews up to five rows.

    Blank rows are skipped; metadata-only rows must still have an identity.
    Integral numeric Excel IDs of up to 15 digits are accepted. String IDs
    must use decimal digits. Missing IDs may be derived from supported URLs.
    ID-only rows use /d2l/home/ID. Excel requires its processed URL header.
    URLs cannot contain credentials, queries, fragments, or foreign origins.
    No hyperlinks are followed, formulas evaluated, or files rewritten.
    """
    source = settings.courses
    try:
        data = source.file.read_bytes()
    except FileNotFoundError:
        raise InventoryError("Course inventory file does not exist.") from None
    except OSError:
        raise InventoryError("Cannot read the course inventory file.") from None
    if source.format == "csv":
        reader = _csv_rows(data)
    elif source.format == "xlsx":
        reader = _excel_rows(data, source.sheet)
    else:
        raise InventoryError("Unsupported course inventory format.")

    courses: dict[tuple[str, str], CourseRef] = {}
    diagnostics: list[InventoryDiagnostic] = []
    failures: list[InventoryDiagnostic] = []
    row_count = 0
    with closing(reader):
        first = next(reader, None)
        if first is None:
            raise InventoryError("Inventory has no header row.")
        _, headers = first
        columns = _columns(headers, source)
        for row, values in reader:
            if all(_blank(value) for value in values):
                continue
            row_count += 1
            try:
                if source.format == "csv" and len(values) > len(headers):
                    raise InventoryError("CSV row has more cells than headers.")
                course = _course(
                    values, columns, settings.instance.tenant_origin, row
                )
            except InventoryError as error:
                failures.append(
                    InventoryDiagnostic(
                        code="INVALID_ROW", row=row, message=str(error)
                    )
                )
                continue
            _deduplicate(courses, course, diagnostics)
    if failures:
        preview = "; ".join(
            f"row {item.row}: {item.message}" for item in failures[:5]
        )
        raise InventoryError(
            f"Inventory has {len(failures)} invalid course row(s). {preview}",
            diagnostics=tuple(failures),
        )
    if not courses:
        raise InventoryError("Inventory contains no course rows.")
    return Inventory(
        courses=tuple(courses.values()),
        source_path=source.file,
        source_format=source.format,
        sheet_name=source.sheet if source.format == "xlsx" else None,
        sha256=hashlib.sha256(data).hexdigest(),
        row_count=row_count,
        diagnostics=tuple(diagnostics),
    )
