"""Verify workbook extraction, preservation, publication, and CLI behavior."""

import os
import shutil
import subprocess
import sys
import unittest
from contextlib import chdir, redirect_stderr, redirect_stdout
from dataclasses import FrozenInstanceError
from io import StringIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.hyperlink import Hyperlink

from gc_stem.d2l.config import load_settings
from gc_stem.d2l.courses import load_courses
from gc_stem.d2l.extract_course_urls import (
    ExtractionError,
    extract_course_urls,
    main,
)


TENANT = "https://brightspace.example.org"


class WorkbookExtractionTests(unittest.TestCase):
    """Use real XLSX round trips and injected output failures."""

    def setUp(self) -> None:
        """Create separate input and output paths for each test."""
        self.temporary = TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name).resolve()
        self.source = self.root / "raw courses.xlsx"
        self.destination = self.root / "processed" / "courses.xlsx"

    def fixture(self, *, with_url: bool = True) -> None:
        """Write links, formulas, styles, and an unrelated hidden sheet."""
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "D2L"
        headers = ["Program", "D2l_course_name"]
        if with_url:
            headers.append("CourseURL")
        sheet.append(headers)
        names = (
            "Linked course", "No link", "Internal link", "Formula link",
            "Empty target", "Hash target",
        )
        for row, name in enumerate(names, start=2):
            sheet.cell(row=row, column=1, value="CS")
            sheet.cell(row=row, column=2, value=name)
            if with_url:
                sheet.cell(row=row, column=3, value=f"keep-row-{row}")
        sheet["B2"].hyperlink = f"{TENANT}/d2l/p/home/123"
        sheet["B2"].font = Font(color="0563C1", underline="single")
        sheet["B4"].hyperlink = Hyperlink(ref="B4", location="Notes!A1")
        sheet["B5"] = '=HYPERLINK("https://example.org", "Formula link")'
        sheet["B6"].hyperlink = Hyperlink(ref="B6", target="")
        sheet["B7"].hyperlink = "#Notes!A1"
        sheet["B1"].fill = PatternFill("solid", fgColor="D9EAF7")
        sheet.column_dimensions["B"].width = 42
        sheet.row_dimensions[3].hidden = True
        sheet.row_dimensions[2].height = 24
        sheet.freeze_panes = "C2"
        if with_url:
            sheet["C2"].number_format = "@"
            sheet["C3"].comment = Comment("Keep this note", "Maintainer")
        notes = workbook.create_sheet("Notes")
        notes.sheet_state = "hidden"
        notes["A1"] = "Unrelated data"
        notes["A2"] = "External reference"
        notes["A2"].hyperlink = "https://example.org/reference"
        notes["B2"] = "=1+1"
        notes.merge_cells("D4:E4")
        notes["D4"] = "Merged note"
        workbook.save(self.source)
        workbook.close()

    def output(self) -> Workbook:
        """Open processed output and register prompt cleanup."""
        workbook = load_workbook(self.destination, data_only=False)
        self.addCleanup(workbook.close)
        return workbook

    def assert_no_temporary_files(self) -> None:
        """Check that a publish attempt leaves no temporary workbook."""
        if self.destination.parent.exists():
            self.assertEqual(list(self.destination.parent.glob(".d2l-*")), [])

    def test_existing_url_column_is_reused(self) -> None:
        """The actual external target is copied into the existing column."""
        self.fixture()
        result = extract_course_urls(self.source, self.destination)
        sheet = self.output()["D2L"]
        self.assertEqual(sheet.max_column, 3)
        self.assertEqual(sheet["C2"].value, f"{TENANT}/d2l/p/home/123")
        self.assertEqual(result.copied_count, 1)
        self.assertEqual(result.rows_examined, 6)
        self.assertEqual(result.skipped_count, 5)
        self.assertFalse(result.url_column_created)

    def test_missing_url_column_is_appended(self) -> None:
        """An unprepared workbook gains a URL header and extracted values."""
        self.fixture(with_url=False)
        result = extract_course_urls(self.source, self.destination)
        sheet = self.output()["D2L"]
        self.assertEqual(sheet["C1"].value, "CourseURL")
        self.assertEqual(sheet["C2"].value, f"{TENANT}/d2l/p/home/123")
        self.assertIsNone(sheet["C3"].value)
        self.assertTrue(result.url_column_created)

    def test_no_target_rows_keep_their_values(self) -> None:
        """Rows without an external target retain their existing URLs."""
        self.fixture()
        extract_course_urls(self.source, self.destination)
        sheet = self.output()["D2L"]
        for row in range(3, 8):
            with self.subTest(row=row):
                self.assertEqual(sheet.cell(row, 3).value, f"keep-row-{row}")
        self.assertEqual(sheet["B5"].data_type, "f")
        self.assertEqual(sheet["B4"].hyperlink.location, "Notes!A1")

    def test_existing_url_formula_without_target_is_preserved(self) -> None:
        """Leaving a URL cell unchanged includes its formula and cell type."""
        self.fixture()
        workbook = load_workbook(self.source)
        workbook["D2L"]["C3"] = '=CONCAT("keep", "-formula")'
        workbook.save(self.source)
        workbook.close()
        extract_course_urls(self.source, self.destination)
        cell = self.output()["D2L"]["C3"]
        self.assertEqual(cell.value, '=CONCAT("keep", "-formula")')
        self.assertEqual(cell.data_type, "f")

    def test_formula_marker_target_is_written_as_text(self) -> None:
        """Stored hyperlink data never becomes an executable output formula."""
        self.fixture()
        workbook = load_workbook(self.source)
        workbook["D2L"]["B2"].hyperlink = "=1+1"
        workbook.save(self.source)
        workbook.close()
        extract_course_urls(self.source, self.destination)
        cell = self.output()["D2L"]["C2"]
        self.assertEqual(cell.value, "=1+1")
        self.assertEqual(cell.data_type, "s")

    def test_other_sheets_links_styles_and_layout_survive(self) -> None:
        """Representative openpyxl-supported features survive the round trip."""
        self.fixture()
        extract_course_urls(self.source, self.destination)
        workbook = self.output()
        self.assertEqual(workbook.sheetnames, ["D2L", "Notes"])
        self.assertEqual(workbook["Notes"].sheet_state, "hidden")
        self.assertEqual(workbook["Notes"]["A1"].value, "Unrelated data")
        self.assertEqual(workbook["Notes"]["B2"].value, "=1+1")
        self.assertEqual(
            workbook["Notes"]["A2"].hyperlink.target,
            "https://example.org/reference",
        )
        self.assertIn("D4:E4", workbook["Notes"].merged_cells)
        sheet = workbook["D2L"]
        self.assertEqual(
            sheet["B2"].hyperlink.target, f"{TENANT}/d2l/p/home/123"
        )
        self.assertEqual(sheet["B2"].font.color.rgb, "000563C1")
        self.assertEqual(sheet["B2"].font.underline, "single")
        self.assertEqual(sheet["B1"].fill.fgColor.rgb, "00D9EAF7")
        self.assertEqual(sheet["C2"].number_format, "@")
        self.assertEqual(sheet["C3"].comment.text, "Keep this note")
        self.assertEqual(sheet.column_dimensions["B"].width, 42)
        self.assertTrue(sheet.row_dimensions[3].hidden)
        self.assertEqual(sheet.row_dimensions[2].height, 24)
        self.assertEqual(sheet.freeze_panes, "C2")

    def test_source_bytes_and_timestamp_are_unchanged(self) -> None:
        """Extraction writes only the separate processed workbook."""
        self.fixture()
        original = self.source.read_bytes()
        modified = self.source.stat().st_mtime_ns
        extract_course_urls(self.source, self.destination)
        self.assertEqual(self.source.read_bytes(), original)
        self.assertEqual(self.source.stat().st_mtime_ns, modified)
        self.assert_no_temporary_files()

    def test_result_paths_counts_and_immutability(self) -> None:
        """The result describes the published output without mutable state."""
        self.fixture()
        result = extract_course_urls(str(self.source), str(self.destination))
        self.assertEqual(result.source_path, self.source)
        self.assertEqual(result.destination_path, self.destination)
        self.assertEqual(result.sheet_name, "D2L")
        with self.assertRaises(FrozenInstanceError):
            result.copied_count = 99

    def test_custom_worksheet_and_column_names(self) -> None:
        """Explicit mappings support a different workbook layout."""
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Master shells"
        sheet.append([" Title ", " Link "])
        sheet.append(["Course", "old"])
        sheet["A2"].hyperlink = f"{TENANT}/d2l/home/123"
        workbook.save(self.source)
        workbook.close()
        result = extract_course_urls(
            self.source, self.destination, sheet_name="Master shells",
            name_column="title", url_column="LINK",
        )
        self.assertFalse(result.url_column_created)
        self.assertEqual(
            self.output()["Master shells"]["B2"].value,
            f"{TENANT}/d2l/home/123",
        )

    def test_headers_are_trimmed_and_case_normalized(self) -> None:
        """Cosmetic differences do not create a duplicate URL column."""
        self.fixture()
        workbook = load_workbook(self.source)
        workbook["D2L"]["B1"] = " coursename "
        workbook["D2L"]["C1"] = " COURSEURL "
        workbook.save(self.source)
        workbook.close()
        result = extract_course_urls(self.source, self.destination)
        self.assertFalse(result.url_column_created)
        self.assertEqual(result.copied_count, 1)

    def test_duplicate_headers_fail_before_output_creation(self) -> None:
        """Ambiguous columns are reported without selecting one silently."""
        self.fixture()
        workbook = load_workbook(self.source)
        workbook["D2L"]["D1"] = " coursename "
        workbook.save(self.source)
        workbook.close()
        with self.assertRaisesRegex(ExtractionError, "Duplicate"):
            extract_course_urls(self.source, self.destination)
        self.assertFalse(self.destination.parent.exists())

    def test_missing_name_header_is_reported(self) -> None:
        """Missing required input does not create a processed workbook."""
        self.fixture()
        workbook = load_workbook(self.source)
        workbook["D2L"]["B1"] = "Different"
        workbook.save(self.source)
        workbook.close()
        with self.assertRaisesRegex(ExtractionError, "course-name"):
            extract_course_urls(self.source, self.destination)
        self.assertFalse(self.destination.exists())

    def test_nontext_header_is_rejected(self) -> None:
        """Unexpected header values are not converted into guessed mappings."""
        self.fixture()
        workbook = load_workbook(self.source)
        workbook["D2L"]["A1"] = 123
        workbook.save(self.source)
        workbook.close()
        with self.assertRaisesRegex(ExtractionError, "headers must contain"):
            extract_course_urls(self.source, self.destination)

    def test_missing_worksheet_is_reported(self) -> None:
        """A different worksheet is not processed as a fallback."""
        self.fixture()
        with self.assertRaisesRegex(ExtractionError, "worksheet is missing"):
            extract_course_urls(
                self.source, self.destination, sheet_name="Not present"
            )
        self.assertFalse(self.destination.exists())

    def test_source_destination_identity_is_always_refused(self) -> None:
        """Overwrite permission never permits replacing the source workbook."""
        self.fixture()
        original = self.source.read_bytes()
        for overwrite in (False, True):
            with self.subTest(overwrite=overwrite):
                with self.assertRaises(ExtractionError):
                    extract_course_urls(
                        self.source, self.source, overwrite=overwrite
                    )
        self.assertEqual(self.source.read_bytes(), original)

    def test_hard_link_alias_of_source_is_refused(self) -> None:
        """Different path names cannot bypass source identity protection."""
        self.fixture()
        alias = self.root / "alias.xlsx"
        try:
            os.link(self.source, alias)
        except OSError:
            self.skipTest("Filesystem does not permit hard links.")
        with self.assertRaisesRegex(ExtractionError, "same file"):
            extract_course_urls(self.source, alias, overwrite=True)

    def test_symbolic_link_alias_of_source_is_refused(self) -> None:
        """Path resolution catches source aliases on supporting platforms."""
        self.fixture()
        alias = self.root / "alias.xlsx"
        try:
            alias.symlink_to(self.source)
        except OSError:
            self.skipTest("Platform does not permit symbolic links.")
        with self.assertRaises(ExtractionError):
            extract_course_urls(self.source, alias, overwrite=True)

    def test_existing_output_requires_explicit_overwrite(self) -> None:
        """A default invocation preserves earlier output byte for byte."""
        self.fixture()
        self.destination.parent.mkdir()
        self.destination.write_bytes(b"previous output")
        with self.assertRaises(FileExistsError):
            extract_course_urls(self.source, self.destination)
        self.assertEqual(self.destination.read_bytes(), b"previous output")
        self.assert_no_temporary_files()

    def test_explicit_overwrite_publishes_the_new_workbook(self) -> None:
        """Requested replacement succeeds only after the new workbook saves."""
        self.fixture()
        self.destination.parent.mkdir()
        self.destination.write_bytes(b"previous output")
        extract_course_urls(self.source, self.destination, overwrite=True)
        self.assertEqual(
            self.output()["D2L"]["C2"].value, f"{TENANT}/d2l/p/home/123"
        )
        self.assert_no_temporary_files()

    def test_save_failure_preserves_output_and_cleans_temp(self) -> None:
        """A partial temporary save cannot replace an earlier workbook."""
        self.fixture()
        self.destination.parent.mkdir()
        self.destination.write_bytes(b"previous output")

        def broken_save(workbook: Workbook, stream: object) -> None:
            """Simulate a disk failure after a partial serialized write."""
            stream.write(b"partial workbook")
            raise OSError("AUTH_SECRET internal save details")

        with patch.object(Workbook, "save", broken_save):
            with self.assertRaises(OSError) as caught:
                extract_course_urls(
                    self.source, self.destination, overwrite=True
                )
        self.assertNotIn("AUTH_SECRET", str(caught.exception))
        self.assertEqual(self.destination.read_bytes(), b"previous output")
        self.assert_no_temporary_files()

    def test_serialization_failure_is_sanitized_and_cleans_temp(self) -> None:
        """Nonfilesystem writer errors also leave no output or leaked values."""
        self.fixture()
        with patch.object(
            Workbook, "save", side_effect=ValueError("AUTH_SECRET")
        ):
            with self.assertRaises(OSError) as caught:
                extract_course_urls(self.source, self.destination)
        self.assertNotIn("AUTH_SECRET", str(caught.exception))
        self.assertFalse(self.destination.exists())
        self.assert_no_temporary_files()

    def test_replacement_failure_preserves_old_output(self) -> None:
        """A locked destination remains intact when replacement fails."""
        self.fixture()
        self.destination.parent.mkdir()
        self.destination.write_bytes(b"previous output")
        with patch(
            "gc_stem.d2l.extract_course_urls.os.replace",
            side_effect=PermissionError("Destination is locked"),
        ):
            with self.assertRaises(PermissionError):
                extract_course_urls(
                    self.source, self.destination, overwrite=True
                )
        self.assertEqual(self.destination.read_bytes(), b"previous output")
        self.assert_no_temporary_files()

    def test_unsupported_publication_never_overwrites(self) -> None:
        """A filesystem without hard links fails without creating output."""
        self.fixture()
        with patch(
            "gc_stem.d2l.extract_course_urls.os.link",
            side_effect=OSError("Hard links are unsupported"),
        ):
            with self.assertRaises(OSError):
                extract_course_urls(self.source, self.destination)
        self.assertFalse(self.destination.exists())
        self.assert_no_temporary_files()

    def test_concurrent_output_is_preserved(self) -> None:
        """A destination created during saving is preserved by default."""
        self.fixture()
        real_link = os.link

        def competing_output(source: Path, destination: Path) -> None:
            """Create a competing destination just before publication."""
            destination.write_bytes(b"other publisher")
            real_link(source, destination)

        with patch(
            "gc_stem.d2l.extract_course_urls.os.link", competing_output
        ):
            with self.assertRaises(FileExistsError):
                extract_course_urls(self.source, self.destination)
        self.assertEqual(self.destination.read_bytes(), b"other publisher")
        self.assert_no_temporary_files()

    def test_sync_failure_leaves_no_published_output(self) -> None:
        """A failed flush-to-disk check cannot produce a completed output."""
        self.fixture()
        with patch(
            "gc_stem.d2l.extract_course_urls.os.fsync",
            side_effect=OSError("Sync failed"),
        ):
            with self.assertRaises(OSError):
                extract_course_urls(self.source, self.destination)
        self.assertFalse(self.destination.exists())
        self.assert_no_temporary_files()

    def test_missing_input_is_reported(self) -> None:
        """A nonexistent workbook fails before creating any output paths."""
        with self.assertRaises(FileNotFoundError):
            extract_course_urls(self.source, self.destination)
        self.assertFalse(self.destination.parent.exists())

    def test_corrupt_input_is_reported_without_contents(self) -> None:
        """Parser exceptions cannot copy input contents into diagnostics."""
        self.source.write_bytes(b"AUTH_SECRET is not an XLSX archive")
        with self.assertRaises(ExtractionError) as caught:
            extract_course_urls(self.source, self.destination)
        self.assertNotIn("AUTH_SECRET", str(caught.exception))
        self.assertTrue(caught.exception.__suppress_context__)
        self.assertFalse(self.destination.parent.exists())

    def test_directory_paths_are_rejected(self) -> None:
        """Directories cannot be treated as input or output workbooks."""
        self.source.mkdir()
        with self.assertRaises(ExtractionError):
            extract_course_urls(self.source, self.destination)
        self.source.rmdir()
        self.fixture()
        self.destination.mkdir(parents=True)
        with self.assertRaises(ExtractionError):
            extract_course_urls(self.source, self.destination, overwrite=True)

    def test_only_xlsx_source_and_destination_are_supported(self) -> None:
        """Unsupported file formats cannot silently lose macro data."""
        self.fixture()
        for source, destination in (
            (self.source.with_suffix(".xlsm"), self.destination),
            (self.source.with_suffix(".xls"), self.destination),
            (self.source, self.destination.with_suffix(".xlsm")),
        ):
            with self.subTest(source=source, destination=destination):
                with self.assertRaisesRegex(ExtractionError, ".xlsx"):
                    extract_course_urls(source, destination)

    def test_invalid_arguments_are_rejected(self) -> None:
        """Invalid flags, names, and overlapping column roles fail clearly."""
        self.fixture()
        for options in (
            {"overwrite": "false"},
            {"sheet_name": " "},
            {"name_column": ""},
            {"url_column": None},
            {"url_column": " coursename "},
        ):
            with self.subTest(options=options):
                with self.assertRaises(ExtractionError):
                    extract_course_urls(
                        self.source, self.destination, **options
                    )

    def test_empty_and_invalid_paths_are_rejected(self) -> None:
        """Empty paths cannot resolve silently to a working directory."""
        self.fixture()
        for value in ("", " ", "\x00"):
            with self.subTest(value=value):
                with self.assertRaises(ExtractionError):
                    extract_course_urls(value, self.destination)
                with self.assertRaises(ExtractionError):
                    extract_course_urls(self.source, value)

    def test_merged_url_cells_fail_without_publishing(self) -> None:
        """An unwritable output cell cannot produce a partial workbook."""
        self.fixture()
        workbook = load_workbook(self.source)
        workbook["D2L"]["B3"].hyperlink = f"{TENANT}/d2l/home/456"
        workbook["D2L"].merge_cells("C2:C3")
        workbook.save(self.source)
        workbook.close()
        with self.assertRaisesRegex(ExtractionError, "Unmerge"):
            extract_course_urls(self.source, self.destination)
        self.assertFalse(self.destination.exists())

    def test_relative_paths_use_the_invocation_directory(self) -> None:
        """The utility does not infer input paths from its source checkout."""
        self.fixture()
        with chdir(self.root):
            result = extract_course_urls(
                self.source.name, Path("processed") / "courses.xlsx"
            )
        self.assertEqual(result.source_path, self.source)
        self.assertEqual(result.destination_path, self.destination)

    def test_cli_reports_copied_count_and_success(self) -> None:
        """The standalone command reports the published workbook."""
        self.fixture()
        output = StringIO()
        with redirect_stdout(output):
            status = main([str(self.source), str(self.destination)])
        self.assertEqual(status, 0)
        self.assertIn("Saved 1 course URLs", output.getvalue())
        self.assertIn(str(self.destination), output.getvalue())

    def test_cli_input_failure_returns_two(self) -> None:
        """Expected input failures are readable and do not print tracebacks."""
        errors = StringIO()
        with redirect_stderr(errors):
            status = main([str(self.source), str(self.destination)])
        self.assertEqual(status, 2)
        self.assertIn("Source workbook does not exist", errors.getvalue())
        self.assertNotIn("Traceback", errors.getvalue())

    def test_cli_output_io_failure_returns_one(self) -> None:
        """The CLI distinguishes output I/O failures from invalid input."""
        self.fixture()
        errors = StringIO()
        with patch(
            "gc_stem.d2l.extract_course_urls.os.link",
            side_effect=PermissionError("Publication denied"),
        ):
            with redirect_stderr(errors):
                status = main([str(self.source), str(self.destination)])
        self.assertEqual(status, 1)
        self.assertIn("Publication denied", errors.getvalue())
        self.assertNotIn("Traceback", errors.getvalue())

    def test_module_runs_outside_a_repository_checkout(self) -> None:
        """Import and module invocation work without a nearby pyproject.toml."""
        self.fixture()
        package_root = Path(__file__).resolve().parents[2]
        isolated = self.root / "installed package"
        module = isolated / "gc_stem/d2l/extract_course_urls.py"
        module.parent.mkdir(parents=True)
        (isolated / "gc_stem/__init__.py").write_text("", encoding="utf-8")
        (module.parent / "__init__.py").write_text("", encoding="utf-8")
        shutil.copyfile(
            package_root / "src/gc_stem/d2l/extract_course_urls.py", module
        )
        environment = dict(os.environ)
        environment["PYTHONPATH"] = str(isolated)
        result = subprocess.run(
            [
                sys.executable, "-m", "gc_stem.d2l.extract_course_urls",
                str(self.source), str(self.destination),
            ],
            cwd=self.root, env=environment, capture_output=True,
            text=True, check=False, timeout=30,
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("Saved 1 course URLs", result.stdout)
        self.assertEqual(result.stderr, "")

    def test_processed_output_loads_through_course_inventory_api(self) -> None:
        """Workbook preparation and the pushed loader agree end to end."""
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "D2L"
        sheet.append(["D2l_course_name"])
        for row, identifier in enumerate(("123", "456"), start=2):
            sheet.cell(row, 1, f"Course {identifier}").hyperlink = (
                f"{TENANT}/d2l/p/home/{identifier}"
            )
        workbook.save(self.source)
        workbook.close()
        result = extract_course_urls(self.source, self.destination)
        config = self.root / "d2l.toml"
        config.write_text("schema_version = 1\n", encoding="utf-8")
        settings = load_settings(
            config, environment={}, overrides={
                "instance.base_url": TENANT,
                "courses.file": self.destination,
                "search.terms": ["target"],
                "output.directory": self.root / "reports",
            },
        )
        inventory = load_courses(settings)
        self.assertEqual(result.copied_count, 2)
        self.assertEqual(
            tuple(course.org_unit_id for course in inventory.courses),
            ("123", "456"),
        )


if __name__ == "__main__":
    unittest.main()
